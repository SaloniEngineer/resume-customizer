import io
import pytest
from fastapi.testclient import TestClient
from docx import Document
import fitz

from app.main import app

client = TestClient(app)

def create_sample_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Jane Doe\njane.doe@example.com | (555) 123-4567 | San Francisco, CA\n\nPROFESSIONAL SUMMARY\nSoftware Engineer with 4 years experience building web applications using Python, FastAPI, and React.\n\nEXPERIENCE\nSoftware Engineer | Tech Corp (2021 - Present)\n- Engineered REST APIs in Python and FastAPI, increasing throughput by 25%.\n- Collaborated with frontend team using React and TypeScript.\n\nEDUCATION\nB.S. in Computer Science, University of California (2020)\n")
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes

def create_sample_docx() -> bytes:
    doc = Document()
    doc.add_heading("John Smith", level=1)
    doc.add_paragraph("john.smith@example.com | (555) 987-6543 | New York, NY")
    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph("Full stack developer specializing in Python, Docker, React, and cloud applications.")
    doc.add_heading("Work Experience", level=2)
    doc.add_paragraph("Senior Engineer | Cloud Dynamics (2022 - Present)")
    doc.add_paragraph("• Architected scalable microservices using FastAPI and Docker.")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("B.S. Computer Engineering, MIT, 2021")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def test_health_check():
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ok"
    assert "llm_configured" in data

def test_empty_job_description():
    pdf_bytes = create_sample_pdf()
    response = client.post(
        "/api/customize-resume",
        data={"job_description": ""},
        files={"file": ("resume.pdf", pdf_bytes, "application/pdf")}
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()

def test_short_job_description():
    pdf_bytes = create_sample_pdf()
    response = client.post(
        "/api/customize-resume",
        data={"job_description": "We need a python dev."},
        files={"file": ("resume.pdf", pdf_bytes, "application/pdf")}
    )
    assert response.status_code == 400
    assert "too short" in response.json()["detail"].lower()

def test_invalid_file_extension():
    response = client.post(
        "/api/customize-resume",
        data={"job_description": "Looking for a full stack developer with experience in Python and React. Requirements include 3+ years experience, FastAPI, Docker, and CI/CD pipelines."},
        files={"file": ("resume.exe", b"binarycontent", "application/octet-stream")}
    )
    assert response.status_code == 400
    assert "unsupported" in response.json()["detail"].lower()

def test_customize_resume_pdf():
    pdf_bytes = create_sample_pdf()
    jd = "We are seeking a Full Stack Software Engineer to build scalable backend services in Python and FastAPI, with modern React frontends. Knowledge of Docker, CI/CD, and REST APIs is highly desired."
    
    response = client.post(
        "/api/customize-resume",
        data={"job_description": jd, "generate_cover_letter": "true"},
        files={"file": ("jane_doe_resume.pdf", pdf_bytes, "application/pdf")}
    )
    assert response.status_code == 200
    data = response.json()
    assert "full_name" in data
    assert "professional_summary" in data
    assert "technical_skills" in data
    assert len(data["work_experience"]) > 0
    assert "match_score" in data
    assert "ats_compatibility_score" in data
    assert data["cover_letter"] is not None

def test_customize_resume_docx():
    docx_bytes = create_sample_docx()
    jd = "Looking for a Senior Software Engineer with strong experience in Docker, Python, FastAPI, and Cloud infrastructure. Must have proven microservices background."
    
    response = client.post(
        "/api/customize-resume",
        data={"job_description": jd},
        files={"file": ("john_smith_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert response.status_code == 200
    data = response.json()
    assert "full_name" in data
    assert "work_experience" in data

def test_download_txt():
    # First get a customized resume
    pdf_bytes = create_sample_pdf()
    jd = "Software engineer required with Python, FastAPI and React expertise for full-time role."
    customized = client.post(
        "/api/customize-resume",
        data={"job_description": jd},
        files={"file": ("resume.pdf", pdf_bytes, "application/pdf")}
    ).json()
    
    # Download as TXT
    res = client.post("/api/download/txt", json=customized)
    assert res.status_code == 200
    assert "text/plain" in res.headers["content-type"]
    assert "PROFESSIONAL SUMMARY" in res.text

def test_download_pdf():
    pdf_bytes = create_sample_pdf()
    jd = "Software engineer required with Python, FastAPI and React expertise for full-time role."
    customized = client.post(
        "/api/customize-resume",
        data={"job_description": jd},
        files={"file": ("resume.pdf", pdf_bytes, "application/pdf")}
    ).json()
    
    # Download as PDF
    res = client.post("/api/download/pdf", json=customized)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF")
