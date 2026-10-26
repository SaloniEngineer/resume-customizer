import io
import os
import re
from typing import Tuple, Dict, Any
from fastapi import UploadFile, HTTPException, status
import fitz  # PyMuPDF
from docx import Document

from app.core.config import settings

class ResumeParserService:
    @staticmethod
    async def validate_file(file: UploadFile) -> Tuple[str, bytes]:
        """
        Validates file existence, extension, and size.
        Returns the lowercase extension and file byte content.
        """
        if not file or not file.filename:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="No file provided or filename is missing."
            )
        
        filename = file.filename.strip()
        parts = filename.rsplit(".", 1)
        if len(parts) < 2:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded file lacks a valid extension (.pdf or .docx)."
            )
        
        ext = parts[1].lower()
        if ext not in settings.ALLOWED_EXTENSIONS:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported file type '.{ext}'. Please upload a PDF (.pdf) or Word document (.docx)."
            )
        
        content = await file.read()
        file_size = len(content)
        
        if file_size == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded file is empty (0 bytes). Please upload a valid resume."
            )
        
        if file_size > settings.MAX_FILE_SIZE_BYTES:
            max_mb = settings.MAX_FILE_SIZE_BYTES // (1024 * 1024)
            actual_mb = file_size / (1024 * 1024)
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail=f"File exceeds maximum allowed size of {max_mb} MB (Uploaded: {actual_mb:.2f} MB)."
            )
        
        return ext, content

    @staticmethod
    def extract_text_from_pdf(content: bytes) -> str:
        """
        Extracts textual content from PDF bytes using PyMuPDF (fitz).
        Detects scanned/image-only PDFs and corrupted files.
        """
        try:
            doc = fitz.open(stream=content, filetype="pdf")
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Corrupted or invalid PDF file. Could not read document: {str(e)}"
            )
        
        if doc.page_count == 0:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="The PDF contains no pages."
            )
        
        full_text_parts = []
        has_images_only = True
        
        for page_idx in range(doc.page_count):
            page = doc[page_idx]
            text = page.get_text("text")
            if text and text.strip():
                has_images_only = False
                full_text_parts.append(text.strip())
        
        doc.close()
        
        extracted_text = "\n\n".join(full_text_parts).strip()
        
        # Scanned PDF check
        if not extracted_text or len(extracted_text) < 30:
            if has_images_only:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail="This PDF appears to be a scanned image or contains no selectable text layer. Please provide a text-based PDF or DOCX file."
                )
            else:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail="Extracted text is too short or empty. Please ensure your resume contains readable text."
                )
        
        return ResumeParserService.clean_text(extracted_text)

    @staticmethod
    def extract_text_from_docx(content: bytes) -> str:
        """
        Extracts textual content from DOCX bytes using python-docx.
        """
        try:
            doc = Document(io.BytesIO(content))
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Corrupted or invalid DOCX document: {str(e)}"
            )
        
        full_text_parts = []
        
        # Paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                full_text_parts.append(text)
        
        # Tables (many resumes use tables for layouts)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    full_text_parts.append(" | ".join(row_cells))
        
        extracted_text = "\n".join(full_text_parts).strip()
        
        if not extracted_text or len(extracted_text) < 30:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="The DOCX document contains insufficient or empty text. Please verify the document content."
            )
        
        return ResumeParserService.clean_text(extracted_text)

    @staticmethod
    def clean_text(text: str) -> str:
        """
        Normalizes whitespace and removes odd binary control characters.
        """
        # Remove null bytes and non-printable control chars except tabs/newlines
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", "", text)
        # Normalize redundant blank lines
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @classmethod
    async def parse_resume(cls, file: UploadFile) -> Tuple[str, Dict[str, Any]]:
        """
        Validates and parses the uploaded resume file.
        Returns (extracted_text, metadata).
        """
        ext, content = await cls.validate_file(file)
        
        if ext == "pdf":
            text = cls.extract_text_from_pdf(content)
        elif ext == "docx":
            text = cls.extract_text_from_docx(content)
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported format: .{ext}"
            )
        
        metadata = {
            "filename": file.filename,
            "file_size": len(content),
            "file_type": ext,
            "character_count": len(text),
            "word_count": len(text.split()),
        }
        
        return text, metadata
