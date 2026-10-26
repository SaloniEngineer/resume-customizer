import io
from typing import Dict, Any
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    HRFlowable,
    ListFlowable,
    ListItem
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY

from app.schemas.resume import CustomizedResume

class ResumeGeneratorService:
    @staticmethod
    def generate_txt(resume: CustomizedResume) -> str:
        """
        Generates a clean, ATS-friendly plain text representation of the resume.
        """
        lines = []
        
        # Name
        lines.append(resume.full_name.upper())
        
        # Contact details
        contacts = []
        if resume.contact_info:
            for key, val in resume.contact_info.items():
                if val:
                    contacts.append(str(val))
        if contacts:
            lines.append(" | ".join(contacts))
        
        lines.append("=" * 70)
        lines.append("")
        
        # Professional Summary
        if resume.professional_summary:
            lines.append("PROFESSIONAL SUMMARY")
            lines.append("-" * 30)
            lines.append(resume.professional_summary)
            lines.append("")
        
        # Technical Skills
        if resume.technical_skills:
            lines.append("TECHNICAL SKILLS")
            lines.append("-" * 30)
            for category, skills in resume.technical_skills.items():
                if skills:
                    lines.append(f"• {category}: {', '.join(skills)}")
            lines.append("")
        
        # Work Experience
        if resume.work_experience:
            lines.append("WORK EXPERIENCE")
            lines.append("-" * 30)
            for exp in resume.work_experience:
                header = f"{exp.role} | {exp.company}"
                meta = []
                if exp.duration:
                    meta.append(exp.duration)
                if exp.location:
                    meta.append(exp.location)
                if meta:
                    header += f" ({', '.join(meta)})"
                lines.append(header)
                for bullet in exp.bullets:
                    lines.append(f"  * {bullet}")
                lines.append("")
        
        # Projects
        if resume.projects:
            lines.append("PROJECTS")
            lines.append("-" * 30)
            for proj in resume.projects:
                p_header = proj.name
                if proj.role_or_tech:
                    p_header += f" | {proj.role_or_tech}"
                if proj.link:
                    p_header += f" ({proj.link})"
                lines.append(p_header)
                if proj.description:
                    lines.append(f"  {proj.description}")
                for bullet in proj.bullets:
                    lines.append(f"  * {bullet}")
                lines.append("")
        
        # Education
        if resume.education:
            lines.append("EDUCATION")
            lines.append("-" * 30)
            for edu in resume.education:
                edu_line = f"{edu.degree} - {edu.institution}"
                if edu.graduation_year:
                    edu_line += f" ({edu.graduation_year})"
                if edu.gpa_or_honors:
                    edu_line += f" [{edu.gpa_or_honors}]"
                lines.append(edu_line)
            lines.append("")
        
        # Certifications
        if resume.certifications:
            lines.append("CERTIFICATIONS")
            lines.append("-" * 30)
            for cert in resume.certifications:
                c_line = cert.name
                if cert.issuer:
                    c_line += f" - {cert.issuer}"
                if cert.year:
                    c_line += f" ({cert.year})"
                lines.append(f"• {c_line}")
            lines.append("")
            
        return "\n".join(lines).strip()

    @staticmethod
    def generate_pdf(resume: CustomizedResume) -> bytes:
        """
        Renders a polished, professional PDF resume using ReportLab.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()

        # Custom Palette
        primary_color = colors.HexColor("#1e293b")   # Slate 800
        accent_color = colors.HexColor("#0f766e")    # Teal 700
        text_color = colors.HexColor("#334155")      # Slate 700
        divider_color = colors.HexColor("#cbd5e1")   # Slate 300

        name_style = ParagraphStyle(
            'NameStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=20,
            leading=24,
            textColor=primary_color,
            alignment=TA_CENTER,
            spaceAfter=4
        )

        contact_style = ParagraphStyle(
            'ContactStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=text_color,
            alignment=TA_CENTER,
            spaceAfter=12
        )

        section_heading = ParagraphStyle(
            'SectionHeading',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=accent_color,
            spaceBefore=8,
            spaceAfter=4,
            keepWithNext=True
        )

        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13.5,
            textColor=text_color,
            alignment=TA_JUSTIFY,
            spaceAfter=6
        )

        item_header_style = ParagraphStyle(
            'ItemHeader',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=10,
            leading=13,
            textColor=primary_color,
            spaceBefore=4,
            spaceAfter=2,
            keepWithNext=True
        )

        bullet_style = ParagraphStyle(
            'Bullet',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12.5,
            textColor=text_color,
            leftIndent=12,
            spaceAfter=2
        )

        story = []

        # Header: Name
        story.append(Paragraph(resume.full_name, name_style))

        # Contact Info
        contact_items = []
        if resume.contact_info:
            for k, v in resume.contact_info.items():
                if v:
                    contact_items.append(str(v))
        if contact_items:
            story.append(Paragraph(" • ".join(contact_items), contact_style))

        story.append(HRFlowable(width="100%", thickness=1, color=divider_color, spaceAfter=8))

        # Professional Summary
        if resume.professional_summary:
            story.append(Paragraph("PROFESSIONAL SUMMARY", section_heading))
            story.append(Paragraph(resume.professional_summary, body_style))

        # Technical Skills
        if resume.technical_skills:
            story.append(Paragraph("TECHNICAL SKILLS", section_heading))
            for cat, skills in resume.technical_skills.items():
                if skills:
                    txt = f"<b>{cat}:</b> {', '.join(skills)}"
                    story.append(Paragraph(txt, bullet_style))
            story.append(Spacer(1, 4))

        # Work Experience
        if resume.work_experience:
            story.append(Paragraph("WORK EXPERIENCE", section_heading))
            for exp in resume.work_experience:
                meta = []
                if exp.duration:
                    meta.append(exp.duration)
                if exp.location:
                    meta.append(exp.location)
                meta_str = f" <i>({', '.join(meta)})</i>" if meta else ""
                
                header_text = f"<b>{exp.role}</b> — {exp.company}{meta_str}"
                story.append(Paragraph(header_text, item_header_style))
                for bullet in exp.bullets:
                    story.append(Paragraph(f"• {bullet}", bullet_style))
                story.append(Spacer(1, 4))

        # Projects
        if resume.projects:
            story.append(Paragraph("PROJECTS", section_heading))
            for proj in resume.projects:
                p_text = f"<b>{proj.name}</b>"
                if proj.role_or_tech:
                    p_text += f" | <i>{proj.role_or_tech}</i>"
                if proj.link:
                    p_text += f" [<font color='navy'>{proj.link}</font>]"
                story.append(Paragraph(p_text, item_header_style))
                if proj.description:
                    story.append(Paragraph(proj.description, body_style))
                for b in proj.bullets:
                    story.append(Paragraph(f"• {b}", bullet_style))
                story.append(Spacer(1, 4))

        # Education
        if resume.education:
            story.append(Paragraph("EDUCATION", section_heading))
            for edu in resume.education:
                line = f"<b>{edu.degree}</b>, {edu.institution}"
                if edu.graduation_year:
                    line += f" ({edu.graduation_year})"
                if edu.gpa_or_honors:
                    line += f" — <i>{edu.gpa_or_honors}</i>"
                story.append(Paragraph(line, body_style))

        # Certifications
        if resume.certifications:
            story.append(Paragraph("CERTIFICATIONS", section_heading))
            for cert in resume.certifications:
                line = f"• <b>{cert.name}</b>"
                if cert.issuer:
                    line += f" — {cert.issuer}"
                if cert.year:
                    line += f" ({cert.year})"
                story.append(Paragraph(line, bullet_style))

        doc.build(story)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    @staticmethod
    def generate_docx(resume: CustomizedResume) -> bytes:
        """
        Renders a cleanly formatted Microsoft Word document (.docx).
        """
        doc = Document()
        
        # Set margins to 0.7 inches
        for section in doc.sections:
            section.top_margin = Inches(0.7)
            section.bottom_margin = Inches(0.7)
            section.left_margin = Inches(0.7)
            section.right_margin = Inches(0.7)

        # Name
        name_p = doc.add_paragraph()
        name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = name_p.add_run(resume.full_name)
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(30, 41, 59)

        # Contact info
        contact_items = []
        if resume.contact_info:
            for k, v in resume.contact_info.items():
                if v:
                    contact_items.append(str(v))
        if contact_items:
            contact_p = doc.add_paragraph()
            contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c_run = contact_p.add_run(" | ".join(contact_items))
            c_run.font.size = Pt(9.5)
            c_run.font.color.rgb = RGBColor(100, 116, 139)

        def add_heading(title: str):
            h = doc.add_heading(level=2)
            run = h.add_run(title)
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(15, 118, 110)

        # Summary
        if resume.professional_summary:
            add_heading("PROFESSIONAL SUMMARY")
            p = doc.add_paragraph(resume.professional_summary)
            p.paragraph_format.space_after = Pt(6)

        # Skills
        if resume.technical_skills:
            add_heading("TECHNICAL SKILLS")
            for cat, skills in resume.technical_skills.items():
                if skills:
                    p = doc.add_paragraph(style='List Bullet')
                    r_cat = p.add_run(f"{cat}: ")
                    r_cat.bold = True
                    p.add_run(", ".join(skills))

        # Experience
        if resume.work_experience:
            add_heading("WORK EXPERIENCE")
            for exp in resume.work_experience:
                header_p = doc.add_paragraph()
                r_role = header_p.add_run(f"{exp.role}")
                r_role.bold = True
                header_p.add_run(f" | {exp.company}")
                meta = []
                if exp.duration:
                    meta.append(exp.duration)
                if exp.location:
                    meta.append(exp.location)
                if meta:
                    r_meta = header_p.add_run(f"  ({', '.join(meta)})")
                    r_meta.italic = True
                
                for bullet in exp.bullets:
                    bp = doc.add_paragraph(bullet, style='List Bullet')
                    bp.paragraph_format.space_after = Pt(2)

        # Projects
        if resume.projects:
            add_heading("PROJECTS")
            for proj in resume.projects:
                p_head = doc.add_paragraph()
                r_name = p_head.add_run(proj.name)
                r_name.bold = True
                if proj.role_or_tech:
                    p_head.add_run(f" | {proj.role_or_tech}")
                if proj.description:
                    doc.add_paragraph(proj.description)
                for b in proj.bullets:
                    doc.add_paragraph(b, style='List Bullet')

        # Education
        if resume.education:
            add_heading("EDUCATION")
            for edu in resume.education:
                ep = doc.add_paragraph()
                r_deg = ep.add_run(edu.degree)
                r_deg.bold = True
                ep.add_run(f", {edu.institution}")
                if edu.graduation_year:
                    ep.add_run(f" ({edu.graduation_year})")

        # Certifications
        if resume.certifications:
            add_heading("CERTIFICATIONS")
            for cert in resume.certifications:
                cp = doc.add_paragraph(style='List Bullet')
                r_c = cp.add_run(cert.name)
                r_c.bold = True
                if cert.issuer:
                    cp.add_run(f" — {cert.issuer}")
                if cert.year:
                    cp.add_run(f" ({cert.year})")

        out = io.BytesIO()
        doc.save(out)
        out.seek(0)
        return out.getvalue()
